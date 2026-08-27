"""Build the PRI-Audit psoriasis manuscript from the supplied reference DOCX.

The reference is the design authority: A4 portrait, 1.25-inch side margins,
Times New Roman hierarchy, single-column flow, compact tables, figures in text,
and declarations followed by numbered references. The reference's review marks
and retinal-study content are deliberately removed.
"""

from __future__ import annotations

import csv
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / ".tmp_docx_template.87VD3w" / "reference.docx"
SOURCE = ROOT / "manuscript" / "PRI_Audit_psoriasis_manuscript_v1.md"
OUTPUT = ROOT / "manuscript" / "PRI_Audit_psoriasis_manuscript_v1.docx"
FIGURE_ROOT = ROOT / "figures" / "paper_v14" / "main"
TABLE_ROOT = ROOT / "results" / "paper_v14"

BLACK = RGBColor(0, 0, 0)
MID_GRAY = "D9D9D9"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
FONT = "Times New Roman"


FIGURES = {
    "FIGURE_1": (
        "fig1_pri_audit_framework.png",
        "Figure 1. PRI-Audit defines patient-level resolution-value testing across psoriasis transcriptomes. "
        "(a) Patients or donors are the independent inferential and resampling units; samples, cells and spots "
        "remain nested measurements. (b) Eleven public datasets comprise seven bulk disease-direction cohorts, "
        "one single-cell development cohort, one paired external single-cell cohort, one paired spatial cohort "
        "and one perturbation prior. (c) In strict outer-patient evaluation, the held-out patient is excluded "
        "from feature selection, scaling, representation fitting and model fitting. (d) Candidate methods are "
        "compared with matched expression references. (e) Resolution-specific falsification uses label "
        "permutation, cross-patient cell reassignment and within-sample coordinate shuffle.",
        "PRI-Audit framework showing nested patients, samples, cells or spots, strict outer fitting, matched "
        "incremental value and resolution-specific falsification.",
    ),
    "FIGURE_2": (
        "fig2_expression_ceiling.png",
        "Figure 2. Patient-level expression establishes a high cross-resolution performance ceiling. "
        "(a) Held-cohort gene-direction cosine across seven bulk LS/NL cohorts; labels give paired-patient counts. "
        "(b) Strict outer-fold expression AUROC in GSE151177, GSE173706 and GSE206391; error bars are 95% "
        "patient/donor bootstrap intervals. (c) PCA and NMF paired ΔAUROC relative to strict pseudobulk "
        "expression. (d) Range and mean increment by representation family. Solid vertical lines denote zero; "
        "dashed lines denote the prespecified +0.05 practical-value threshold.",
        "Four-panel benchmark of bulk expression direction, strict expression AUROC, paired low-rank increments "
        "and representation-family gain ranges.",
    ),
    "FIGURE_3": (
        "fig3_unsupervised_optimism.png",
        "Figure 3. Whole-dataset unsupervised fitting produces measurable optimism. (a) Global and strict "
        "representation-fitting boundaries. (b) Global versus fold-refitted neighborhood AUROC for K=8, 16 and "
        "32. (c) Sample-level bootstrap distributions of the paired global-minus-strict gap from 5,000 draws. "
        "(d) Sensitivity to equal-patient versus cell-count weighting. (e) Full distribution of 500 cross-patient "
        "cell shuffles preserving sample cell counts. The orange line is observed global-neighborhood AUROC; "
        "finite-permutation P=0.006.",
        "Five-panel audit showing global versus strict neighborhood fitting, optimism bootstrap, weighting "
        "sensitivity and the full cross-patient cell-shuffle null distribution.",
    ),
    "FIGURE_4": (
        "fig4_external_transfer.png",
        "Figure 4. Fixed cell-neighborhood transfer preserves ranking but not calibration or incremental value. "
        "(a) GSE151177 PCA, KMeans and neighborhoods were frozen and applied to GSE173706 without target "
        "refitting. (b) ROC across 11 complete PN/PP donor pairs (AUROC 0.884). (c) Paired donor probabilities. "
        "(d) Five-bin quantile reliability curve; calibration coefficients are descriptive and unstable. "
        "(e) Paired donor-bootstrap ΔAUROC relative to strict expression: −0.116 (95% CI −0.281 to −0.008). "
        "(f) Full 2,000-draw paired-label permutation null; finite-permutation P=0.0055.",
        "Six-panel external-transfer figure containing the frozen mapping, ROC, paired probabilities, reliability "
        "curve, paired incremental value and paired-label null distribution.",
    ),
    "FIGURE_5": (
        "fig5_strict_spatial_value.png",
        "Figure 5. Strict spatial graph statistics add no patient-level value beyond matched expression. "
        "(a) Per-gene spatial features based on six-neighbor Moran-like agreement and local contrast. "
        "(b) Predefined representative LS/NL sections, colored by the fixed S100A7/S100A7A/KRT16/IL36G/CXCL8/"
        "SERPINB4 module. (c) Strict patient-LODO AUROC for expression-500, spatial graph and their combination. "
        "(d) Patient-level mean LS-minus-NL prediction difference for all 12 patients. (e) Paired increments "
        "relative to expression-500. (f) Full 100-draw within-sample coordinate-shuffle null; P=0.446.",
        "Six-panel spatial benchmark showing graph definitions, representative spatial maps, matched AUROCs, "
        "paired patient scores, incremental values and coordinate-shuffle null distribution.",
    ),
    "FIGURE_6": (
        "fig6_falsification_evidence_map.png",
        "Figure 6. Falsification separates reproducible disease signal from representation-specific gain. "
        "(a) Nine candidate-versus-matched-expression ΔAUROC estimates and 95% patient/donor bootstrap "
        "intervals. (b) Observed-minus-null AUROC under cell, label and coordinate falsification. Cell and label "
        "disruption remove disease-associated information, whereas coordinate disruption has approximately zero "
        "effect on strict spatial-graph AUROC. (c) Evidence-grade matrix distinguishing satisfied, failed and "
        "non-applicable criteria. Patient-level disease signal is supported; reproducible resolution-specific "
        "incremental value is not demonstrated.",
        "Evidence summary containing all paired increments, observed-minus-null falsification effects and the "
        "final evidence-grade matrix.",
    ),
}


EQUATIONS = {
    "EQUATION_1": "rᵢ,ₘ⁽⁻ᵍ⁾ = Tₘ(Xᵢ; θ̂ₘ⁽⁻ᵍ⁾),    θ̂ₘ⁽⁻ᵍ⁾ = Fit({Xⱼ: j ∉ g}).",
    "EQUATION_2": "Δₘ⁽ᵈ⁾ = AUROC(y, p̂candidate,m⁽ᵈ⁾) − AUROC(y, p̂reference,m⁽ᵈ⁾).",
    "EQUATION_3": "πₘ(δ₀) = Bootstrap Pr(Δₘ ≥ δ₀),    δ₀ = 0.05.",
    "EQUATION_4": "Oₖ = AUROC(global, K) − AUROC(strict, K).",
    "EQUATION_5": "logit{Pr(yᵢ = 1 | p̂ᵢ)} = α + β logit(p̂ᵢ).",
    "EQUATION_6": "Iᵢg = [nᵢ⁻¹ Σᵤ zᵢᵤg |Nᵢ(u)|⁻¹ Σᵥ∈Nᵢ(u) zᵢᵥg] / [nᵢ⁻¹ Σᵤ zᵢᵤg² + ε].",
    "EQUATION_7": "Cᵢg = nᵢ⁻¹ Σᵤ |xᵢᵤg − |Nᵢ(u)|⁻¹ Σᵥ∈Nᵢ(u) xᵢᵥg|.",
    "EQUATION_8": "Gₘ = 1[Δₘ,external ≥ 0.05] × 1[CI stable] × 1[strict outer fit] × 1[falsification passed].",
}


def set_run_font(run, size: float | None = None, *, bold=None, italic=None, color=BLACK) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = color
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "en-US")


def get_or_add_style(doc: Document, name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, style_type)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.15

    title = get_or_add_style(doc, "PRI Title")
    title.base_style = normal
    title.font.name = FONT
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True

    for name, size, before, after in [
        ("Heading 1", 16, 12, 6),
        ("Heading 2", 14, 9, 4),
        ("Heading 3", 12, 7, 3),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    abstract = get_or_add_style(doc, "PRI Abstract")
    abstract.base_style = normal
    abstract.font.name = FONT
    abstract.font.size = Pt(12)
    abstract.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract.paragraph_format.space_after = Pt(4)
    abstract.paragraph_format.line_spacing = 1.08

    caption = get_or_add_style(doc, "PRI Figure Caption")
    caption.base_style = normal
    caption.font.name = FONT
    caption.font.size = Pt(10)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.keep_together = True

    table_caption = get_or_add_style(doc, "PRI Table Caption")
    table_caption.base_style = normal
    table_caption.font.name = FONT
    table_caption.font.size = Pt(10)
    table_caption.font.bold = True
    table_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_caption.paragraph_format.space_before = Pt(7)
    table_caption.paragraph_format.space_after = Pt(3)
    table_caption.paragraph_format.keep_with_next = True
    table_caption.paragraph_format.keep_together = True

    equation = get_or_add_style(doc, "PRI Equation")
    equation.base_style = normal
    equation.font.name = FONT
    equation.font.size = Pt(11)
    equation.font.italic = True
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(3)
    equation.paragraph_format.space_after = Pt(7)
    equation.paragraph_format.keep_together = True

    reference = get_or_add_style(doc, "PRI Reference")
    reference.base_style = normal
    reference.font.name = FONT
    reference.font.size = Pt(10)
    reference.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    reference.paragraph_format.left_indent = Inches(0.22)
    reference.paragraph_format.first_line_indent = Inches(-0.22)
    reference.paragraph_format.space_after = Pt(3)
    reference.paragraph_format.line_spacing = 1.0

    table_text = get_or_add_style(doc, "PRI Table Text")
    table_text.base_style = normal
    table_text.font.name = FONT
    table_text.font.size = Pt(8)
    table_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.0


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clear_story(story) -> None:
    element = story._element
    for child in list(element):
        element.remove(child)
    p = OxmlElement("w:p")
    element.append(p)


def configure_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.header_distance = Inches(0.45)
        section.footer_distance = Inches(0.45)
        clear_story(section.header)
        clear_story(section.footer)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(0)
        footer.paragraph_format.space_after = Pt(0)
        run = footer.add_run()
        set_run_font(run, 9)
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        for node in (fld_begin, instr, fld_sep, text, fld_end):
            run._r.append(node)


def add_inline_markdown(paragraph, text: str, size: float = 12) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if is_bold else part
        run = paragraph.add_run(content)
        set_run_font(run, size, bold=is_bold)


def add_body_paragraph(doc: Document, text: str, *, abstract: bool = False):
    style = "PRI Abstract" if abstract else "Normal"
    p = doc.add_paragraph(style=style)
    add_inline_markdown(p, text, 12)
    p.paragraph_format.widow_control = True
    return p


def add_equation(doc: Document, key: str) -> None:
    p = doc.add_paragraph(style="PRI Equation")
    r = p.add_run(EQUATIONS[key])
    set_run_font(r, 11, italic=True)


def set_picture_alt_text(paragraph, title: str, description: str) -> None:
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("name", title)
        doc_pr.set("title", title)
        doc_pr.set("descr", description)


def add_figure(doc: Document, key: str) -> None:
    filename, caption, alt = FIGURES[key]
    path = FIGURE_ROOT / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(5.72))
    set_picture_alt_text(p, key.replace("_", " ").title(), alt)
    cap = doc.add_paragraph(style="PRI Figure Caption")
    add_inline_markdown(cap, caption, 10)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge, size, color in (
        ("top", "8", "000000"),
        ("left", "4", "A6A6A6"),
        ("bottom", "8", "000000"),
        ("right", "4", "A6A6A6"),
        ("insideH", "4", "BFBFBF"),
        ("insideV", "4", "D9D9D9"),
    ):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(table)


def mark_header_repeat(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def write_cell(cell, value: str, *, header=False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.style = "PRI Table Text"
    p.alignment = align
    p.paragraph_format.keep_together = True
    segments = str(value).split("\n")
    for i, segment in enumerate(segments):
        if i:
            p.add_run().add_break(WD_BREAK.LINE)
        r = p.add_run(segment)
        set_run_font(r, 8, bold=header)
    if header:
        set_cell_shading(cell, MID_GRAY)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def add_table_1(doc: Document) -> None:
    cap = doc.add_paragraph(style="PRI Table Caption")
    add_inline_markdown(cap, "Table 1. Public datasets and independent analysis units.", 10)
    rows = read_csv(TABLE_ROOT / "main_table1_public_datasets.csv")
    headers = ["Dataset", "Modality", "Design", "Independent unit", "n", "Analytical role"]
    widths = [0.72, 0.68, 1.05, 0.92, 0.48, 1.90]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        write_cell(cell, value, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    mark_header_repeat(table.rows[0])
    for idx, row in enumerate(rows):
        n = row["independent_units_after_filtering"]
        n = str(int(float(n))) if n else "—"
        values = [
            row["dataset"],
            row["modality"],
            row["design"],
            row["independent_unit"],
            n,
            row["role"],
        ]
        cells = table.add_row().cells
        for col, (cell, value) in enumerate(zip(cells, values)):
            if idx % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            write_cell(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if col in (0, 1, 4) else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_geometry(table, widths)
    note = doc.add_paragraph(style="PRI Figure Caption")
    add_inline_markdown(note, "Note: Cells and spots quantify measurement density and do not increase the number of independent inferential units.", 9)


def fmt_num(value: str, digits=3) -> str:
    return f"{float(value):.{digits}f}".replace("-", "−")


def add_table_2(doc: Document) -> None:
    cap = doc.add_paragraph(style="PRI Table Caption")
    add_inline_markdown(cap, "Table 2. Patient-level incremental value and evidence grade.", 10)
    rows = read_csv(TABLE_ROOT / "main_table2_patient_level_incremental_value.csv")
    headers = ["Dataset", "Candidate representation", "Candidate AUROC", "Reference AUROC", "Paired ΔAUROC [95% CI]", "Evidence"]
    widths = [0.70, 1.28, 0.62, 0.62, 1.50, 1.03]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        write_cell(cell, value, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    mark_header_repeat(table.rows[0])
    for idx, row in enumerate(rows):
        delta = f"{fmt_num(row['delta_auroc'])}\n[{fmt_num(row['delta_ci95_low'])}, {fmt_num(row['delta_ci95_high'])}]"
        evidence = f"P(Δ ≥ 0.05)={float(row['probability_delta_ge_0_05']):.3f}\nNo increment"
        values = [
            row["dataset"],
            row["candidate_representation"],
            fmt_num(row["candidate_auroc"]),
            fmt_num(row["reference_auroc"]),
            delta,
            evidence,
        ]
        cells = table.add_row().cells
        for col, (cell, value) in enumerate(zip(cells, values)):
            if idx % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            write_cell(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if col != 1 else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_geometry(table, widths)
    note = doc.add_paragraph(style="PRI Figure Caption")
    add_inline_markdown(note, "The paired bootstrap resampled patients or donors. The prespecified practical-value threshold was ΔAUROC ≥ 0.05.", 9)


def add_frontmatter(doc: Document, lines: list[str]) -> int:
    title = doc.add_paragraph(style="PRI Title")
    r = title.add_run(lines[0][2:].strip())
    set_run_font(r, 16, bold=True)

    author = doc.add_paragraph()
    author.paragraph_format.space_after = Pt(2)
    author.paragraph_format.keep_with_next = True
    r = author.add_run(lines[2])
    set_run_font(r, 12)

    affiliation = doc.add_paragraph()
    affiliation.paragraph_format.space_after = Pt(2)
    affiliation.paragraph_format.keep_with_next = True
    r = affiliation.add_run(lines[4])
    set_run_font(r, 10.5)

    corr = doc.add_paragraph()
    corr.paragraph_format.space_after = Pt(8)
    corr.paragraph_format.keep_with_next = True
    r = corr.add_run(lines[6])
    set_run_font(r, 10.5)
    return 8


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)
    configure_page(doc)
    configure_styles(doc)

    props = doc.core_properties
    props.title = "Falsification-first patient-level benchmarking reveals limited incremental value of single-cell and spatial representations in psoriasis"
    props.subject = "PRI-Audit psoriasis transcriptomics manuscript draft"
    props.author = ""
    props.keywords = "psoriasis; patient-level validation; single-cell RNA-seq; spatial transcriptomics; falsification"
    props.comments = "Initial manuscript draft generated from the frozen v14 evidence package."

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = add_frontmatter(doc, lines)
    in_abstract = False
    in_references = False

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        index += 1
        if not line:
            continue
        token = re.fullmatch(r"\[\[(FIGURE_\d|TABLE_\d|EQUATION_\d)\]\]", line)
        if token:
            key = token.group(1)
            if key.startswith("FIGURE"):
                add_figure(doc, key)
            elif key == "TABLE_1":
                add_table_1(doc)
            elif key == "TABLE_2":
                add_table_2(doc)
            else:
                add_equation(doc, key)
            continue
        if line.startswith("## "):
            heading = line[3:]
            in_abstract = heading == "Abstract"
            in_references = heading == "References"
            p = doc.add_heading(heading, level=1)
            for run in p.runs:
                set_run_font(run, 16, bold=True)
            continue
        if line.startswith("### "):
            p = doc.add_heading(line[4:], level=2)
            for run in p.runs:
                set_run_font(run, 14, bold=True)
            continue
        if in_references and re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="PRI Reference")
            add_inline_markdown(p, line, 10)
            continue
        add_body_paragraph(doc, line, abstract=in_abstract)

    # Keep the title block and Abstract together as much as practical.
    for p in doc.paragraphs[:5]:
        p.paragraph_format.keep_with_next = True

    # Avoid stale automatic fields from the reference and request refresh of PAGE.
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
