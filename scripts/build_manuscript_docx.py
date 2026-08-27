"""Build a rendered-review DOCX from the frozen manuscript and result tables.

Design preset: narrative_proposal (US Letter; 1 in margins; Calibri 11;
justified, 1.333 body spacing; blue heading hierarchy; fixed-width tables).
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript" / "psoriasis_frozen_representation_manuscript.md"
OUT = ROOT / "manuscript" / "psoriasis_frozen_representation_manuscript.docx"
FIG = ROOT / "figures" / "paper"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT = "F4F6F9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "120")
    ind.set(qn("w:type"), "dxa")
    tbl_pr.append(ind)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def plain_text(md: str) -> str:
    return re.sub(r"\*\*(.*?)\*\*", r"\1", md).replace("*", "")


def add_markdown_paragraph(doc: Document, text: str, style: str = "Normal") -> None:
    para = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        run = para.add_run(plain_text(part))
        run.bold = part.startswith("**") and part.endswith("**")


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.333
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.base_style = doc.styles["Normal"]
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_text = doc.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    table_text.base_style = doc.styles["Normal"]
    table_text.font.size = Pt(8)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.0


def add_dataframe(doc: Document, title: str, frame: pd.DataFrame, widths: list[float]) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(4)
    run = para.add_run(title)
    run.bold = True
    run.font.size = Pt(10)
    table = doc.add_table(rows=1, cols=len(frame.columns), style="Table Grid")
    headers = [str(x).replace("_", " ") for x in frame.columns]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        set_cell_shading(cell, LIGHT)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    set_table_geometry(table, widths)


def add_figures(doc: Document) -> None:
    captions = [
        ("fig1_study_design_and_boundaries.png", "Figure 1. Study design and leakage boundaries. Mechanism discovery, disease-model development and locking, technical blind transport, and treatment-direction analyses used non-overlapping roles. The final claim is recovery tracking, not diagnostic superiority."),
        ("fig2_mechanism_residual_boundary.png", "Figure 2. The proposed synergy residual was weakly reproducible across experiments. Gene-level residual effects in GSE24767 and GSE86451 had Spearman rho=0.21; the bootstrap distribution is centered near 0.19."),
        ("fig3_development_algorithm_boundary.png", "Figure 3. Development analysis falsified the proposed anchor advantage. Mean LOCO AUROC peaked at gamma=1 (ordinary ridge); the residual coefficient was not sign-stable across held-out cohorts."),
        ("fig4_diagnostic_transport_boundary.png", "Figure 4. Diagnostic transport boundary in GSE223799. The locked ridge, v0 NNLS and NOS2−CCL27 each perfectly separated five PSO from five AD samples at both macro and micro scales."),
        ("fig5_treatment_direction.png", "Figure 5. Treatment-associated direction of the locked score. GSE226244 is sample-level baseline versus week 12; GSE85034 ADA and MTX estimates are paired patient-level week 16 versus LS."),
        ("fig6_nonlesional_normalization_controls.png", "Figure 6. Fixed-score non-lesional normalization control. All three fixed scores moved toward each patient's own non-lesional state by week 16. This is a post-hoc descriptive control; raw score scales are not comparable."),
    ]
    doc.add_heading("Figures", level=1)
    for image, caption in captions:
        doc.add_picture(str(FIG / image), width=Inches(6.25))
        p = doc.add_paragraph(caption, style="Figure Caption")
        p.paragraph_format.keep_together = True


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    configure_styles(doc)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("Public-data psoriasis transcriptomics | manuscript draft").font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Draft for scientific review — not for clinical use").font.size = Pt(8)

    lines = MANUSCRIPT.read_text().splitlines()
    skip_tables = False
    for line in lines:
        if line.startswith("## Tables"):
            skip_tables = True
            continue
        if line.startswith("## References"):
            skip_tables = False
        if skip_tables:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            r = p.add_run(line[2:])
            r.font.name = "Calibri"
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = RGBColor.from_string("0B2545")
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(plain_text(line))
            r.bold = True
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.add_run(line)
        elif line.strip():
            add_markdown_paragraph(doc, line)

    doc.add_page_break()
    add_figures(doc)
    doc.add_page_break()
    doc.add_heading("Tables", level=1)
    registry = pd.read_csv(ROOT / "results" / "paper_dataset_registry.csv")
    compact_registry = registry.assign(
        role=registry.role.str.replace("mechanism discovery", "mechanism", regex=False)
        .str.replace("development disease cohort", "development", regex=False)
        .str.replace("independent treatment replication", "treatment", regex=False)
    )[["accession", "role", "assay", "status"]]
    add_dataframe(doc, "Table 1. Dataset registry", compact_registry, [1.0, 2.45, 1.1, 1.95])

    doc.add_page_break()
    table2 = pd.DataFrame([
        ["Development LOCO", "n=120", "AUROC 0.906", "gamma=1, lambda=1; anchor did not win"],
        ["GSE223799 macro", "5 PSO / 5 AD", "all three AUROC 1.000", "comparison saturated"],
        ["GSE223799 micro", "5 PSO / 5 AD", "all three AUROC 1.000", "comparison saturated"],
        ["GSE261704", "not analyzed", "8.0% coverage", "pre-specified exclusion"],
        ["GSE153007", "not analyzed", "non-lesional samples", "not a lesion comparison"],
    ], columns=["analysis", "unit", "estimate", "boundary"])
    add_dataframe(doc, "Table 2. Model development and external transport", table2, [1.45, 1.3, 1.55, 2.2])

    doc.add_page_break()
    table3 = pd.DataFrame([
        ["GSE226244", "week 12 - baseline", "unpaired samples", "-0.606", "[-0.713, -0.491]", "group-level only"],
        ["GSE85034 ADA", "week 16 - LS", "14 paired patients", "-0.800", "[-0.956, -0.631]", "treatment-associated direction"],
        ["GSE85034 MTX", "week 16 - LS", "14 paired patients", "-0.437", "[-0.761, -0.098]", "treatment-associated direction"],
        ["GSE85034 pooled", "week 16 - LS", "28 paired patients", "-0.618", "[-0.800, -0.414]", "25/28 decreased"],
    ], columns=["dataset", "contrast", "unit", "mean delta", "95% CI", "interpretation"])
    add_dataframe(doc, "Table 3. Treatment-direction results", table3, [1.05, 1.1, 1.25, 0.75, 1.15, 1.2])

    doc.add_page_break()
    control = pd.DataFrame([
        ["locked ridge", 28, "-0.382", "[-0.571, -0.193]", "0.9999"],
        ["v0 NNLS", 28, "-0.010", "[-0.015, -0.005]", "1.0000"],
        ["NOS2-CCL27", 28, "-0.811", "[-1.133, -0.497]", "1.0000"],
    ], columns=["fixed score", "n", "distance change", "95% CI", "P(change < 0)"])
    add_dataframe(doc, "Supplementary Table S1. Week-16 distance to patient's own non-lesional sample", control, [1.5, 0.55, 1.3, 1.7, 1.45])
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
