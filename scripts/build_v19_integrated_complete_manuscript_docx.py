"""Build the PRI-Audit psoriasis manuscript from the supplied reference DOCX.

The reference is the design authority: A4 portrait, 1.25-inch side margins,
Times New Roman hierarchy, single-column flow, compact tables, figures in text,
and declarations followed by numbered references. The reference's review marks
and retinal-study content are deliberately removed.
"""

from __future__ import annotations

import csv
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / ".tmp_docx_template.87VD3w" / "reference.docx"
SOURCE = ROOT / "manuscript" / "PRI_Audit_psoriasis_manuscript_v6_integrated_complete.md"
OUTPUT = ROOT / "manuscript" / "PRI_Audit_psoriasis_manuscript_v6_integrated_complete.docx"
FIGURE_ROOT = ROOT / "figures" / "paper_v18" / "main"
ED_FIGURE_ROOT = ROOT / "figures" / "paper_v18" / "extended_data"
TABLE_ROOT = ROOT / "results" / "paper_v15"

BLACK = RGBColor(0, 0, 0)
MID_GRAY = "D9D9D9"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
FONT = "Times New Roman"


FIGURES = {
    "FIGURE_1": (
        "fig1_pri_audit_technical_route.png",
        "Figure 1. PRI-Audit technical route for the public psoriasis benchmark. Stage A assigns fixed roles to "
        "seven paired bulk cohorts, two single-cell cohorts, one spatial cohort and one perturbation prior. "
        "Stage B constructs fold-local resolution encoders and a matched expression reference before computing "
        "paired out-of-fold ΔAUROC. Stage C performs patient- or donor-level resampling, frozen external transfer, "
        "strict spatial validation and information-source-specific null tests. Stage D admits only bounded claims "
        "that satisfy patient isolation, strict outer fitting, stable uncertainty and targeted falsification.",
        "Four-stage technical route from public-data foundation through PRI-Audit construction and strict "
        "validation to evidence-gated outputs.",
    ),
    "FIGURE_2": (
        "fig2_pri_audit_module_architecture.png",
        "Figure 2. PRI-Audit algorithm architecture and module functions. The patient or donor defines the outer "
        "fold, bootstrap unit and permissible inference. M1 locks the fold registry; M2 fits adaptive preprocessing "
        "using training groups; M3 produces expression, low-rank, cell-state, spatial-graph and exploratory "
        "trigger-transport sample vectors; M4 compares candidate and matched-reference predictions; M5 produces "
        "complete label, cell and coordinate null distributions; and M6 maps the resulting evidence to supported, "
        "bounded or refuted claim language. The held-out group bypasses every fitting operation.",
        "Six-module architecture with independent-unit boundary, encoder library, auditable handoffs, targeted "
        "falsification and evidence gate.",
    ),
    "FIGURE_3": (
        "fig3_public_data_landscape.png",
        "Figure 3. Public-data landscape for the patient-level psoriasis audit. Seven paired bulk cohorts show "
        "cohort medians and interquartile ranges of patient-level LS-minus-NL rank shifts for fixed epidermal, "
        "IL-17/innate and interferon-associated programs. GSE151177 and GSE173706 show measured cells in the "
        "fixed source coordinate system; GSE206391 shows observed LS/NL tissue spots from the preselected "
        "median-effect patient; and GSE24767 shows across-donor mean cytokine-response trajectories. Complete "
        "patient and donor distributions are retained in Extended Data Figures 13 and 14. These portraits are "
        "descriptive and did not enter model fitting, performance estimation or claim selection.",
        "Low-saturation landscape of seven paired bulk cohorts, two single-cell cohorts, one spatial cohort and "
        "one keratinocyte perturbation cohort.",
    ),
    "FIGURE_4": (
        "fig4_expression_ceiling.png",
        "Figure 4. Patient-level expression establishes a high cross-resolution performance ceiling. "
        "(a) Held-cohort gene-direction cosine across seven bulk LS/NL cohorts; labels give paired-patient counts. "
        "(b) Strict outer-fold expression AUROC in GSE151177, GSE173706 and GSE206391; error bars are 95% "
        "patient/donor bootstrap intervals. (c) PCA and NMF paired ΔAUROC relative to strict pseudobulk "
        "expression. (d) Range and mean increment by representation family. Solid vertical lines denote zero; "
        "dashed lines denote the prespecified +0.05 practical-value threshold.",
        "Four-panel benchmark of bulk expression direction, strict expression AUROC, paired low-rank increments "
        "and representation-family gain ranges.",
    ),
    "FIGURE_5": (
        "fig5_unsupervised_optimism.png",
        "Figure 5. Whole-dataset unsupervised fitting produces measurable optimism. (a) Global and strict "
        "representation-fitting boundaries. (b) Global versus fold-refitted neighborhood AUROC for K=8, 16 and "
        "32. (c) Sample-level bootstrap distributions of the paired global-minus-strict gap from 5,000 draws. "
        "(d) Sensitivity to equal-patient versus cell-count weighting. (e) Full distribution of 500 cross-patient "
        "cell shuffles preserving sample cell counts. The orange line is observed global-neighborhood AUROC; "
        "finite-permutation P=0.006.",
        "Five-panel audit showing global versus strict neighborhood fitting, optimism bootstrap, weighting "
        "sensitivity and the full cross-patient cell-shuffle null distribution.",
    ),
    "FIGURE_6": (
        "fig6_external_transfer.png",
        "Figure 6. Fixed cell-neighborhood transfer preserves ranking but not calibration or incremental value. "
        "(a) GSE151177 PCA, KMeans and neighborhoods were frozen and applied to GSE173706 without target "
        "refitting. (b) ROC across 11 complete PN/PP donor pairs (AUROC 0.884). (c) Paired donor probabilities. "
        "(d) Five-bin quantile reliability curve; calibration coefficients are descriptive and unstable. "
        "(e) Paired donor-bootstrap ΔAUROC relative to strict expression: −0.116 (95% CI −0.281 to −0.008). "
        "(f) Full 2,000-draw paired-label permutation null; finite-permutation P=0.0055.",
        "Six-panel external-transfer figure containing the frozen mapping, ROC, paired probabilities, reliability "
        "curve, paired incremental value and paired-label null distribution.",
    ),
    "FIGURE_7": (
        "fig7_strict_spatial_value.png",
        "Figure 7. Strict spatial graph statistics add no patient-level value beyond matched expression. "
        "(a) Per-gene spatial features based on six-neighbor Moran-like agreement and local contrast. "
        "(b) Predefined representative LS/NL sections, colored by the fixed S100A7/S100A7A/KRT16/IL36G/CXCL8/"
        "SERPINB4 module. (c) Strict patient-LODO AUROC for expression-500, spatial graph and their combination. "
        "(d) Patient-level mean LS-minus-NL prediction difference for all 12 patients. (e) Paired increments "
        "relative to expression-500. (f) Full 100-draw within-sample coordinate-shuffle null; P=0.446.",
        "Six-panel spatial benchmark showing graph definitions, representative spatial maps, matched AUROCs, "
        "paired patient scores, incremental values and coordinate-shuffle null distribution.",
    ),
    "FIGURE_8": (
        "fig8_falsification_evidence_map.png",
        "Figure 8. Falsification separates reproducible disease signal from representation-specific gain. "
        "(a) Nine candidate-versus-matched-expression ΔAUROC estimates and 95% patient/donor bootstrap "
        "intervals. (b) Observed-minus-null AUROC under cell, label and coordinate falsification. Cell and label "
        "disruption remove disease-associated information, whereas coordinate disruption has approximately zero "
        "effect on strict spatial-graph AUROC. (c) Evidence-grade matrix distinguishing satisfied, failed and "
        "non-applicable criteria. Patient-level disease signal is supported; reproducible resolution-specific "
        "incremental value is not demonstrated.",
        "Evidence summary containing all paired increments, observed-minus-null falsification effects and the "
        "final evidence-grade matrix.",
    ),
    "FIGURE_9": (
        "fig9_observed_data_atlas.png",
        "Figure 9. Observed single-cell and spatial landscapes reveal cohort shift and patient heterogeneity. "
        "(a,b) Fixed-source UMAP of 19,868 GSE151177 cells colored by condition and the 16 frozen source "
        "neighborhoods. (c,d) A display subset of 14,493 GSE173706 cells transformed through the frozen source "
        "scaler, PCA, UMAP and neighborhood definitions without target refitting. (e) GSE206391 NL/LS sections "
        "from patients nearest the predefined low, median and high quantiles of the strict paired combined-model "
        "effect. Color denotes the fixed S100A7/S100A7A/KRT16/IL36G/CXCL8/SERPINB4 module. The maps are "
        "descriptive and did not enter performance estimation or model selection.",
        "Data-derived atlas of source and external single-cell landscapes plus six predefined spatial sections "
        "from low-, median- and high-effect patients.",
    ),
    "FIGURE_10": (
        "fig10_robustness_multiverse.png",
        "Figure 10. Resampling and multiverse analyses preserve the bounded conclusion. (a) Strict expression "
        "ROC curves in the three high-resolution datasets. (b) All nine candidate-versus-matched-expression "
        "paired ΔAUROC estimates with patient/donor-bootstrap 95% intervals; the dashed line denotes +0.05. "
        "(c) Mean strict spatial-graph AUROC by HVG and neighbor count across regularization and seeds. (d) "
        "Regularization-by-seed surface for the primary 500-HVG, six-neighbor specification. (e) Cell, external "
        "paired-label and strict coordinate null distributions; diamonds denote observed AUROC. Absolute "
        "hyperparameter variation does not establish matched incremental value.",
        "Five-panel robustness summary containing strict ROC curves, the full incremental-value forest, spatial "
        "sensitivity heatmaps and resolution-specific null distributions.",
    ),
}


INTEGRATED_FIGURES = {
    "FIGURE_1": (
        ROOT / "figures" / "paper_v19_fig1_dataset_atlas.png",
        "Figure 1. Eleven public psoriasis datasets represented by data-derived views. Seven paired bulk cohorts "
        "show patient-level LS-minus-NL shifts for the same fixed epidermal, IL-17/innate and interferon programs. "
        "The high-resolution panels show observed GSE151177 and GSE173706 single-cell coordinates, paired "
        "GSE206391 spatial sections and donor-resolved GSE24767 cytokine responses. The 191 bulk pairs, 34,361 "
        "displayed single cells, 12 spatial patients and three perturbation donors remain linked to their biological "
        "units. These panels are descriptive and were not used for feature, sample or claim selection.",
        "Data-derived overview of all eleven public psoriasis datasets across bulk, single-cell, spatial and "
        "keratinocyte perturbation scales.",
    ),
    "FIGURE_2": (
        ROOT / "figures" / "paper_v19_fig2_overall_framework.png",
        "Figure 2. Data-integrated technical route for PRI-Audit. Stage A assigns fixed roles to all 11 datasets "
        "while retaining a representative data view for each cohort. Stage B enforces patient/donor isolation and "
        "fits adaptive objects within outer training folds. Stage C compares each candidate with matched expression "
        "on the same held-out units and estimates patient/donor bootstrap uncertainty. Stage D applies cell, label "
        "and coordinate nulls before the claim gate. The final output distinguishes reproducible disease-state signal "
        "from unsupported resolution-specific incremental value.",
        "Four-stage PRI-Audit framework containing all eleven real dataset views, fold-local construction, matched "
        "audit and falsification-gated claims.",
    ),
    "FIGURE_3": (
        ROOT / "figures" / "paper_v19_fig3_algorithm_architecture.png",
        "Figure 3. PRI-Audit algorithm architecture and module functions. Real bulk, single-cell, spatial and "
        "perturbation panels enter the model through four modality lanes. M1 defines the patient/donor registry; M2 "
        "fits feature selection, scaling and representation objects only on outer-fold training units; M3 constructs "
        "matched expression, low-rank, cell-neighborhood, spatial-graph and exploratory trigger-transport vectors; "
        "M4 compares candidate and expression predictions under the same folds, feature budget and classifier; M5 "
        "produces patient/donor intervals, calibration and information-source-specific nulls; and M6 maps the evidence "
        "to supported, bounded or refuted language. Held-out units bypass all fitting operations.",
        "Six-module PRI-Audit model architecture driven by real bulk, UMAP, spatial and perturbation data views.",
    ),
    "FIGURE_8": (
        ROOT / "figures" / "paper_v19_fig4_evidence_closure.png",
        "Figure 8. Evidence closure for the primary claim. Strict expression ROC curves establish transferable "
        "patient-level disease signal; global-versus-strict comparisons expose optimism from whole-dataset fitting; "
        "external reliability shows that ranking does not imply calibrated probability; and matched spatial "
        "increments plus the coordinate-shuffle null show no reproducible coordinate-specific gain. Together these "
        "results support a recurrent psoriasis disease-state signal while bounding claims of incremental value from "
        "low-rank, cell-neighborhood and spatial-graph representations.",
        "Four-panel evidence synthesis using strict ROC, optimism, external calibration and spatial coordinate-null "
        "results, followed by the evidence-graded conclusion.",
    ),
}

FIGURES.update(INTEGRATED_FIGURES)


ED_FIGURES = {
    "ED_FIGURE_1": ("figED1_cell_spot_qc.png", "Extended Data Figure 1. Cohort scale and quality control.", "Quality-control summary across the three high-resolution datasets."),
    "ED_FIGURE_2": ("figED2_strict_outer_roc.png", "Extended Data Figure 2. Strict outer-fold ROC curves.", "Strict expression, PCA and NMF ROC curves."),
    "ED_FIGURE_3": ("figED3_prediction_distributions.png", "Extended Data Figure 3. Strict prediction distributions.", "Outer-fold prediction distributions by dataset, method and outcome."),
    "ED_FIGURE_4": ("figED4_bulk_robustness.png", "Extended Data Figure 4. Seven-cohort bulk robustness.", "Cross-cohort bulk expression robustness diagnostics."),
    "ED_FIGURE_5": ("figED5_neighborhood_sensitivity.png", "Extended Data Figure 5. Cell-neighborhood sensitivity.", "Patient weighting, fit-boundary and optimism diagnostics."),
    "ED_FIGURE_6": ("figED6_external_calibration.png", "Extended Data Figure 6. External-transfer calibration diagnostics.", "Reliability, calibration and Brier diagnostics for external transfer."),
    "ED_FIGURE_7": ("figED7_transport_diagnostics.png", "Extended Data Figure 7. Perturbation-informed transport diagnostics.", "Exploratory trigger-informed transport and fused-GW results."),
    "ED_FIGURE_8": ("figED8_spatial_graph_diagnostics.png", "Extended Data Figure 8. Spatial-graph patient-level diagnostics.", "Strict patient-level spatial graph diagnostics."),
    "ED_FIGURE_9": ("figED9_complete_null_distributions.png", "Extended Data Figure 9. Complete resolution-specific null distributions.", "Full cell, label and coordinate null distributions."),
    "ED_FIGURE_10": ("figED10_hyperparameter_sensitivity.png", "Extended Data Figure 10. Hyperparameter sensitivity matrices.", "HVG, neighbor, regularization and seed sensitivity matrices."),
    "ED_FIGURE_11": ("figED11_observed_single_cell_landscapes.png", "Extended Data Figure 11. Observed single-cell landscapes and composition shifts.", "Observed source and target single-cell embeddings plus neighborhood composition shifts."),
    "ED_FIGURE_12": ("figED12_predefined_spatial_gallery.png", "Extended Data Figure 12. Predefined multi-patient spatial gallery.", "Six observed spatial sections selected by a predefined patient-effect quantile rule."),
    "ED_FIGURE_13": ("figED13_bulk_patient_program_distributions.png", "Extended Data Figure 13. Patient-level bulk program distributions.", "Every paired patient is retained as a dot within the fixed epidermal, IL-17/innate and interferon-associated programs; thick segments are interquartile ranges and white-centered markers are medians."),
    "ED_FIGURE_14": ("figED14_calm_cytokine_donor_responses.png", "Extended Data Figure 14. Donor-resolved keratinocyte cytokine responses.", "GSE24767 donor-paired program trajectories for IL-17, TNF, combination and the synergy residual, followed by the across-donor mean and standard deviation."),
}


EQUATIONS = {
    "EQUATION_1": "rᵢ,ₘ⁽⁻ᵍ⁾ = Tₘ(Xᵢ; θ̂ₘ⁽⁻ᵍ⁾),    θ̂ₘ⁽⁻ᵍ⁾ = Fit({Xⱼ: j ∉ g}).",
    "EQUATION_2": "Δₘ⁽ᵈ⁾ = AUROC(y, p̂candidate,m⁽ᵈ⁾) − AUROC(y, p̂reference,m⁽ᵈ⁾).",
    "EQUATION_3": "πₘ(δ₀) = Bootstrap Pr(Δₘ ≥ δ₀),    δ₀ = 0.05.",
    "EQUATION_4": "Oₖ = AUROC(global, K) − AUROC(strict, K).",
    "EQUATION_5": "logit{Pr(yᵢ = 1 | p̂ᵢ)} = α + β logit(p̂ᵢ).",
    "EQUATION_6": "Iᵢg = [nᵢ⁻¹ Σᵤ zᵢᵤg |Nᵢ(u)|⁻¹ Σᵥ∈Nᵢ(u) zᵢᵥg] / [nᵢ⁻¹ Σᵤ zᵢᵤg² + ε].",
    "EQUATION_7": "Cᵢg = nᵢ⁻¹ Σᵤ |xᵢᵤg − |Nᵢ(u)|⁻¹ Σᵥ∈Nᵢ(u) xᵢᵥg|.",
    "EQUATION_8": "Gₘ = 1[Δₘ,external ≥ 0.05] × 1[CI stable] × 1[strict outer fit] × 1[falsification passed].",
}


def set_run_font(run, size: float | None = None, *, bold=None, italic=None, color=BLACK) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = color
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "en-US")


def get_or_add_style(doc: Document, name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, style_type)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.15

    title = get_or_add_style(doc, "PRI Title")
    title.base_style = normal
    title.font.name = FONT
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True

    for name, size, before, after in [
        ("Heading 1", 16, 12, 6),
        ("Heading 2", 14, 9, 4),
        ("Heading 3", 12, 7, 3),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    abstract = get_or_add_style(doc, "PRI Abstract")
    abstract.base_style = normal
    abstract.font.name = FONT
    abstract.font.size = Pt(12)
    abstract.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract.paragraph_format.space_after = Pt(4)
    abstract.paragraph_format.line_spacing = 1.08

    caption = get_or_add_style(doc, "PRI Figure Caption")
    caption.base_style = normal
    caption.font.name = FONT
    caption.font.size = Pt(10)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.keep_together = True

    table_caption = get_or_add_style(doc, "PRI Table Caption")
    table_caption.base_style = normal
    table_caption.font.name = FONT
    table_caption.font.size = Pt(10)
    table_caption.font.bold = True
    table_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_caption.paragraph_format.space_before = Pt(7)
    table_caption.paragraph_format.space_after = Pt(3)
    table_caption.paragraph_format.keep_with_next = True
    table_caption.paragraph_format.keep_together = True

    equation = get_or_add_style(doc, "PRI Equation")
    equation.base_style = normal
    equation.font.name = FONT
    equation.font.size = Pt(11)
    equation.font.italic = True
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(3)
    equation.paragraph_format.space_after = Pt(7)
    equation.paragraph_format.keep_together = True

    reference = get_or_add_style(doc, "PRI Reference")
    reference.base_style = normal
    reference.font.name = FONT
    reference.font.size = Pt(10)
    reference.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    reference.paragraph_format.left_indent = Inches(0.22)
    reference.paragraph_format.first_line_indent = Inches(-0.22)
    reference.paragraph_format.space_after = Pt(3)
    reference.paragraph_format.line_spacing = 1.0

    table_text = get_or_add_style(doc, "PRI Table Text")
    table_text.base_style = normal
    table_text.font.name = FONT
    table_text.font.size = Pt(8)
    table_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.0


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clear_story(story) -> None:
    element = story._element
    for child in list(element):
        element.remove(child)
    p = OxmlElement("w:p")
    element.append(p)


def configure_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.header_distance = Inches(0.45)
        section.footer_distance = Inches(0.45)
        clear_story(section.header)
        clear_story(section.footer)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(0)
        footer.paragraph_format.space_after = Pt(0)
        run = footer.add_run()
        set_run_font(run, 9)
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        for node in (fld_begin, instr, fld_sep, text, fld_end):
            run._r.append(node)


def add_inline_markdown(paragraph, text: str, size: float = 12) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if is_bold else part
        run = paragraph.add_run(content)
        set_run_font(run, size, bold=is_bold)


def add_body_paragraph(doc: Document, text: str, *, abstract: bool = False):
    style = "PRI Abstract" if abstract else "Normal"
    p = doc.add_paragraph(style=style)
    add_inline_markdown(p, text, 12)
    p.paragraph_format.widow_control = True
    return p


def add_equation(doc: Document, key: str) -> None:
    p = doc.add_paragraph(style="PRI Equation")
    r = p.add_run(EQUATIONS[key])
    set_run_font(r, 11, italic=True)


def set_picture_alt_text(paragraph, title: str, description: str) -> None:
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("name", title)
        doc_pr.set("title", title)
        doc_pr.set("descr", description)


def add_figure(doc: Document, key: str) -> None:
    filename, caption, alt = FIGURES[key]
    path = filename if isinstance(filename, Path) else FIGURE_ROOT / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(5.72))
    set_picture_alt_text(p, key.replace("_", " ").title(), alt)
    cap = doc.add_paragraph(style="PRI Figure Caption")
    add_inline_markdown(cap, caption, 10)


def add_extended_figure(doc: Document, key: str) -> None:
    filename, caption, alt = ED_FIGURES[key]
    path = ED_FIGURE_ROOT / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(5.55))
    set_picture_alt_text(p, key.replace("_", " ").title(), alt)
    cap = doc.add_paragraph(style="PRI Figure Caption")
    add_inline_markdown(cap, caption, 10)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge, size, color in (
        ("top", "8", "000000"),
        ("left", "4", "A6A6A6"),
        ("bottom", "8", "000000"),
        ("right", "4", "A6A6A6"),
        ("insideH", "4", "BFBFBF"),
        ("insideV", "4", "D9D9D9"),
    ):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(table)


def mark_header_repeat(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def write_cell(cell, value: str, *, header=False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.style = "PRI Table Text"
    p.alignment = align
    p.paragraph_format.keep_together = True
    segments = str(value).split("\n")
    for i, segment in enumerate(segments):
        if i:
            p.add_run().add_break(WD_BREAK.LINE)
        r = p.add_run(segment)
        set_run_font(r, 8, bold=header)
    if header:
        set_cell_shading(cell, MID_GRAY)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def add_table_1(doc: Document) -> None:
    cap = doc.add_paragraph(style="PRI Table Caption")
    add_inline_markdown(cap, "Table 1. Public datasets and independent analysis units.", 10)
    rows = read_csv(TABLE_ROOT / "main_table1_public_datasets.csv")
    headers = ["Dataset", "Modality", "Design", "Independent unit", "n", "Analytical role"]
    widths = [0.72, 0.68, 1.05, 0.92, 0.48, 1.90]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        write_cell(cell, value, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    mark_header_repeat(table.rows[0])
    for idx, row in enumerate(rows):
        n = row["independent_units_after_filtering"]
        n = str(int(float(n))) if n else "—"
        values = [
            row["dataset"],
            row["modality"],
            row["design"],
            row["independent_unit"],
            n,
            row["role"],
        ]
        cells = table.add_row().cells
        for col, (cell, value) in enumerate(zip(cells, values)):
            if idx % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            write_cell(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if col in (0, 1, 4) else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_geometry(table, widths)
    note = doc.add_paragraph(style="PRI Figure Caption")
    add_inline_markdown(note, "Note: Cells and spots quantify measurement density and do not increase the number of independent inferential units.", 9)


def fmt_num(value: str, digits=3) -> str:
    return f"{float(value):.{digits}f}".replace("-", "−")


def add_table_2(doc: Document) -> None:
    cap = doc.add_paragraph(style="PRI Table Caption")
    add_inline_markdown(cap, "Table 2. Patient-level incremental value and evidence grade.", 10)
    rows = read_csv(TABLE_ROOT / "main_table2_patient_level_incremental_value.csv")
    headers = ["Dataset", "Candidate representation", "Candidate AUROC", "Reference AUROC", "Paired ΔAUROC [95% CI]", "Evidence"]
    widths = [0.70, 1.28, 0.62, 0.62, 1.50, 1.03]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        write_cell(cell, value, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    mark_header_repeat(table.rows[0])
    for idx, row in enumerate(rows):
        delta = f"{fmt_num(row['delta_auroc'])}\n[{fmt_num(row['delta_ci95_low'])}, {fmt_num(row['delta_ci95_high'])}]"
        evidence = f"P(Δ ≥ 0.05)={float(row['probability_delta_ge_0_05']):.3f}\nNo increment"
        values = [
            row["dataset"],
            row["candidate_representation"],
            fmt_num(row["candidate_auroc"]),
            fmt_num(row["reference_auroc"]),
            delta,
            evidence,
        ]
        cells = table.add_row().cells
        for col, (cell, value) in enumerate(zip(cells, values)):
            if idx % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            write_cell(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if col != 1 else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_geometry(table, widths)
    note = doc.add_paragraph(style="PRI Figure Caption")
    add_inline_markdown(note, "The paired bootstrap resampled patients or donors. The prespecified practical-value threshold was ΔAUROC ≥ 0.05.", 9)


def fmt_optional(value: str, digits=3) -> str:
    if value is None or str(value).strip() == "" or str(value).lower() == "nan":
        return "—"
    return fmt_num(str(value), digits)


def add_custom_csv_table(
    doc: Document,
    *,
    caption: str,
    filename: str,
    columns: list[tuple[str, float, object]],
    note: str | None = None,
) -> None:
    cap = doc.add_paragraph(style="PRI Table Caption")
    add_inline_markdown(cap, caption, 10)
    rows = read_csv(TABLE_ROOT / filename)
    table = doc.add_table(rows=1, cols=len(columns))
    for cell, (header, _, _) in zip(table.rows[0].cells, columns):
        write_cell(cell, header, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    mark_header_repeat(table.rows[0])
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index, (cell, (_, _, getter)) in enumerate(zip(cells, columns)):
            value = getter(row)
            if index % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            write_cell(
                cell,
                value,
                align=WD_ALIGN_PARAGRAPH.CENTER if column_index in (0, len(columns) - 2) else WD_ALIGN_PARAGRAPH.LEFT,
            )
    set_table_geometry(table, [width for _, width, _ in columns])
    if note:
        paragraph = doc.add_paragraph(style="PRI Figure Caption")
        add_inline_markdown(paragraph, note, 9)


def add_table_3(doc: Document, *, supplementary=False) -> None:
    add_custom_csv_table(
        doc,
        caption=("Supplementary Table 4. Resolution-specific falsification results." if supplementary else "Table 3. Resolution-specific falsification results."),
        filename=("supp_table4_falsification.csv" if supplementary else "main_table3_resolution_specific_falsification.csv"),
        columns=[
            ("Dataset", 0.65, lambda r: r["dataset"]),
            ("Broken source", 0.93, lambda r: r["information_source_broken"]),
            ("Model", 0.95, lambda r: r["tested_model"]),
            ("Obs.", 0.52, lambda r: fmt_num(r["observed_auroc"])),
            ("Null mean", 0.52, lambda r: fmt_num(r["null_mean_auroc"])),
            ("P", 0.43, lambda r: fmt_num(r["finite_permutation_p"])),
            ("Interpretation", 1.75, lambda r: r["interpretation"]),
        ],
        note="Finite-permutation P values use (1+b)/(B+1). The coordinate null tests coordinate-specific value, not the existence of spatially varying expression.",
    )


def add_table_4(doc: Document) -> None:
    add_custom_csv_table(
        doc,
        caption="Table 4. Diagnostics and sensitivity analyses.",
        filename="main_table4_diagnostics_and_sensitivity.csv",
        columns=[
            ("Domain", 0.88, lambda r: r["domain"]),
            ("Estimand", 1.18, lambda r: r["estimand"]),
            ("Estimate", 0.62, lambda r: fmt_optional(r["estimate"])),
            ("95% interval / range", 1.05, lambda r: r["uncertainty_or_range"].replace("-", "−")),
            ("Unit", 0.88, lambda r: r["analysis_unit"]),
            ("Interpretation", 1.14, lambda r: r["interpretation"]),
        ],
        note="Calibration intervals and optimism intervals use donor- or patient-level bootstrap. The spatial range summarizes the frozen 144-condition multiverse.",
    )


def add_supp_table(doc: Document, number: int) -> None:
    if number == 1:
        add_custom_csv_table(
            doc, caption="Supplementary Table 1. High-resolution dataset quality-control summary.",
            filename="supp_table1_high_resolution_qc.csv",
            columns=[
                ("Dataset", 0.75, lambda r: r["dataset"]),
                ("Independent units", 0.85, lambda r: r["independent_units"]),
                ("Samples / sections", 0.90, lambda r: r["analyzed_samples_or_sections"]),
                ("Raw cells / spots", 1.00, lambda r: f"{int(float(r['raw_cells_or_spots'])):,}"),
                ("Retained", 0.90, lambda r: f"{int(float(r['retained_cells_or_spots'])):,}"),
                ("Retention", 0.75, lambda r: f"{float(r['overall_retention']):.1%}"),
            ],
        )
    elif number == 2:
        add_custom_csv_table(
            doc, caption="Supplementary Table 2. Strict outer-fold resampling results.",
            filename="supp_table2_strict_resampling.csv",
            columns=[
                ("Dataset", 0.72, lambda r: r["dataset"]),
                ("Method", 0.80, lambda r: r["method"]),
                ("n groups", 0.58, lambda r: r["n_groups"]),
                ("AUROC", 0.58, lambda r: fmt_num(r["auroc"])),
                ("95% CI", 1.10, lambda r: f"{fmt_num(r['bootstrap_ci95_low'])} to {fmt_num(r['bootstrap_ci95_high'])}"),
                ("Null mean", 0.68, lambda r: fmt_num(r["permutation_mean"])),
                ("Permutation P", 0.72, lambda r: fmt_num(r["permutation_p"])),
            ],
        )
    elif number == 3:
        add_custom_csv_table(
            doc, caption="Supplementary Table 3. Full paired incremental-value evidence.",
            filename="supp_table3_full_incremental_evidence.csv",
            columns=[
                ("Dataset", 0.62, lambda r: r["dataset"]),
                ("Candidate", 0.96, lambda r: r["candidate_representation"]),
                ("Reference", 0.88, lambda r: r["matched_reference"]),
                ("Δ", 0.50, lambda r: fmt_num(r["delta_auroc"])),
                ("95% CI", 0.85, lambda r: f"{fmt_num(r['delta_ci95_low'])} to {fmt_num(r['delta_ci95_high'])}"),
                ("Pr≥.05", 0.52, lambda r: fmt_num(r["probability_delta_ge_0_05"])),
                ("Status", 0.87, lambda r: r["claim_status"]),
            ],
        )
    elif number == 4:
        add_table_3(doc, supplementary=True)
    elif number == 5:
        add_custom_csv_table(
            doc, caption="Supplementary Table 5. External transfer and calibration metrics.",
            filename="supp_table5_external_transfer_calibration.csv",
            columns=[
                ("Metric", 1.65, lambda r: r["metric"]),
                ("Estimate", 0.95, lambda r: fmt_num(r["estimate"])),
                ("CI low", 0.95, lambda r: fmt_num(r["ci_low"])),
                ("CI high", 0.95, lambda r: fmt_num(r["ci_high"])),
                ("Resampling unit", 1.25, lambda r: r["unit"]),
            ],
        )
    elif number == 6:
        add_custom_csv_table(
            doc, caption="Supplementary Table 6. Spatial sensitivity summary by HVG and neighbor count.",
            filename="supp_table6_spatial_sensitivity_summary.csv",
            columns=[
                ("HVG", 0.80, lambda r: r["hvg"]),
                ("Neighbors", 0.90, lambda r: r["neighbors"]),
                ("Minimum AUROC", 1.20, lambda r: fmt_num(r["min"])),
                ("Mean AUROC", 1.15, lambda r: fmt_num(r["mean"])),
                ("Maximum AUROC", 1.20, lambda r: fmt_num(r["max"])),
            ],
        )
    elif number == 7:
        def shifted_display(row):
            return re.sub(
                r"Fig\. ([1-8])\b",
                lambda match: f"Fig. {int(match.group(1)) + 2}",
                row["primary_display"],
            )

        add_custom_csv_table(
            doc, caption="Supplementary Table 7. Frozen claim registry.",
            filename="supp_table7_claim_registry.csv",
            columns=[
                ("ID", 0.40, lambda r: r["claim_id"]),
                ("Frozen statement", 1.80, lambda r: r["frozen_statement"]),
                ("Evidence", 0.85, lambda r: r["evidence_status"]),
                ("Allowed wording", 1.30, lambda r: r["allowed_wording"]),
                ("Prohibited wording", 1.00, lambda r: r["prohibited_wording"]),
                ("Fig.", 0.40, shifted_display),
            ],
        )


def add_frontmatter(doc: Document, lines: list[str]) -> int:
    title = doc.add_paragraph(style="PRI Title")
    r = title.add_run(lines[0][2:].strip())
    set_run_font(r, 16, bold=True)

    author = doc.add_paragraph()
    author.paragraph_format.space_after = Pt(2)
    author.paragraph_format.keep_with_next = True
    r = author.add_run(lines[2])
    set_run_font(r, 12)

    affiliation = doc.add_paragraph()
    affiliation.paragraph_format.space_after = Pt(2)
    affiliation.paragraph_format.keep_with_next = True
    r = affiliation.add_run(lines[4])
    set_run_font(r, 10.5)

    corr = doc.add_paragraph()
    corr.paragraph_format.space_after = Pt(8)
    corr.paragraph_format.keep_with_next = True
    r = corr.add_run(lines[6])
    set_run_font(r, 10.5)
    return 8


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)
    configure_page(doc)
    configure_styles(doc)

    props = doc.core_properties
    props.title = "Falsification-first patient-level benchmarking reveals limited incremental value of single-cell and spatial representations in psoriasis"
    props.subject = "Complete PRI-Audit psoriasis transcriptomics manuscript with all eleven data-derived views integrated into the framework and model architecture"
    props.author = ""
    props.keywords = "psoriasis; patient-level validation; single-cell RNA-seq; spatial transcriptomics; falsification"
    props.comments = "Integrated complete manuscript generated from the frozen evidence package, the four-slide editable PowerPoint figure suite and explicit PRI-Audit module contracts."

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = add_frontmatter(doc, lines)
    in_abstract = False
    in_references = False
    extended_figure_count = 0
    supplementary_table_count = 0

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        index += 1
        if not line:
            continue
        token = re.fullmatch(
            r"\[\[((?:ED_)?FIGURE_\d+|(?:SUPP_)?TABLE_\d+|EQUATION_\d+)\]\]",
            line,
        )
        if token:
            key = token.group(1)
            if key.startswith("ED_FIGURE"):
                add_extended_figure(doc, key)
            elif key.startswith("FIGURE"):
                add_figure(doc, key)
            elif key == "TABLE_1":
                add_table_1(doc)
            elif key == "TABLE_2":
                add_table_2(doc)
            elif key == "TABLE_3":
                add_table_3(doc)
            elif key == "TABLE_4":
                add_table_4(doc)
            elif key.startswith("SUPP_TABLE_"):
                add_supp_table(doc, int(key.rsplit("_", 1)[1]))
            else:
                add_equation(doc, key)
            continue
        if line.startswith("## "):
            heading = line[3:]
            if heading in {"Extended Data", "Supplementary Tables"}:
                doc.add_page_break()
            in_abstract = heading == "Abstract"
            in_references = heading == "References"
            p = doc.add_heading(heading, level=1)
            for run in p.runs:
                set_run_font(run, 16, bold=True)
            continue
        if line.startswith("### "):
            heading = line[4:]
            if heading.startswith("Extended Data Figure"):
                if extended_figure_count:
                    doc.add_page_break()
                extended_figure_count += 1
            elif heading.startswith("Supplementary Table"):
                if supplementary_table_count:
                    doc.add_page_break()
                supplementary_table_count += 1
            p = doc.add_heading(heading, level=2)
            for run in p.runs:
                set_run_font(run, 14, bold=True)
            continue
        if in_references and re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="PRI Reference")
            add_inline_markdown(p, line, 10)
            continue
        add_body_paragraph(doc, line, abstract=in_abstract)

    # Keep the title block and Abstract together as much as practical.
    for p in doc.paragraphs[:5]:
        p.paragraph_format.keep_with_next = True

    # Avoid stale automatic fields from the reference and request refresh of PAGE.
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
