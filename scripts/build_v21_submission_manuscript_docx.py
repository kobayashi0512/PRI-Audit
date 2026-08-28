"""Build the v21 PRI-Audit submission manuscript from frozen source artifacts.

The retained retinal manuscript is used only as a visual authority (A4,
Times New Roman, compact single-column scientific flow).  Its content and
tracked review artefacts are removed before the PRI-Audit manuscript is built.
"""

from __future__ import annotations

import csv
import os
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REFERENCE = Path(
    "/Users/tongyue/Library/Containers/com.tencent.xinWeChat/Data/Documents/"
    "xwechat_files/wxid_vjuiqasexnx722_174a/msg/file/2026-08/"
    "RetinalFusionNet_4.1.5 (1).docx"
)
REFERENCE = Path(os.environ.get("PRI_AUDIT_REFERENCE_DOCX", str(DEFAULT_REFERENCE)))
SOURCE = ROOT / "manuscript" / "PRI_Audit_psoriasis_manuscript_v8_submission_ready.md"
DEFAULT_OUTPUT = ROOT / "manuscript" / "PRI_Audit_psoriasis_manuscript_v9_submission_ready.docx"
OUTPUT = Path(os.environ.get("PRI_AUDIT_OUTPUT_DOCX", str(DEFAULT_OUTPUT)))
TABLES = ROOT / "results" / "paper_v21"
STYLE_PROFILE = os.environ.get("PRI_AUDIT_STYLE_PROFILE", "submission").strip().lower()
REGIONSUFF_PROFILE = STYLE_PROFILE == "regionsuff"
BMC_PROFILE = STYLE_PROFILE == "bmc"

FONT = "Times New Roman"
BLACK = RGBColor(31, 31, 31)
GOLD = RGBColor(174, 123, 0)
MID_GRAY = "D9D9D9"
LIGHT_GRAY = "F5F5F5"

FIGURES = {
    "FIGURE_1": (
        ROOT / "figures" / "paper_v19_fig1_dataset_atlas.png",
        "Figure 1. Data-derived landscape of the 11 public psoriasis datasets. Seven paired bulk cohorts, two single-cell cohorts, one spatial cohort and one perturbation cohort are displayed at their native biological scales. The panels are descriptive only: they do not enter feature selection, model fitting, performance estimation or claim selection. Cohort roles and independent units are specified in Table 1.",
        "Data-derived atlas of public psoriasis cohorts across bulk, single-cell, spatial and perturbation scales.",
    ),
    "FIGURE_2": (
        ROOT / "figures" / "paper_v20_fig2_overall_framework.png",
        "Figure 2. PRI-Audit framework. Public data are assigned distinct evidential roles; adaptive objects are fitted only in outer-fold training groups; each candidate is compared with matched expression on identical held-out units; and the final claim is filtered through independent-unit uncertainty and representation-specific falsification. The framework supports the bounded conclusion, not a new clinical classifier claim.",
        "Four-stage PRI-Audit workflow from public data to bounded evidence claims.",
    ),
    "FIGURE_3": (
        ROOT / "figures" / "paper_v20_fig3_algorithm_architecture.png",
        "Figure 3. Algorithmic architecture and module contracts. M1 defines the independent unit; M2 fits fold-local preprocessing; M3 constructs candidate and matched-expression representations; M4 computes paired held-out increments; M5 applies source-specific counterfactual diagnostics; and M6 maps only supported evidence to claim language. The held-out unit bypasses every fitting operation.",
        "PRI-Audit module architecture with biological-unit boundary and matched comparator.",
    ),
    "FIGURE_4": (
        ROOT / "figures" / "paper_v21" / "fig11_v21_strictness_and_reference_audit.png",
        "Figure 4. Strictness and matched-reference audit. (A) All strict PCA/NMF increments against matched 2,000-HVG expression and their independent-unit 95% bootstrap intervals. The red dashed line is the pre-specified +0.05 practical-value threshold. (B) In GSE206391, replacing matched expression with a deliberately weakened 16-HVG reference manufactures an apparent PCA gain. (C) Fully nested neighbourhood results show that the global–strict difference is resolution-sensitive rather than universally positive.",
        "Strict matched low-rank comparisons, mismatched-reference sensitivity, and nested-neighbourhood diagnostic.",
    ),
    "FIGURE_5": (
        ROOT / "figures" / "paper_v21" / "fig12_v21_independence_and_spatial_null.png",
        "Figure 5. Independence and spatial-null diagnostics. (A) Random cell/spot splits produce high AUROCs but every test measurement derives from a source sample already represented in training; they estimate within-source measurement interpolation, not new-patient prediction. (B) In GSE206391, 500 within-sample coordinate shuffles retain counts, labels, folds and fold-specific HVGs. Spatial-graph and combined-model null distributions overlap their observed held-out AUROCs, providing no coordinate-specific incremental evidence for the tested graph representation.",
        "Random measurement-split and spatial coordinate-shuffle diagnostics.",
    ),
    "FIGURE_6": (
        ROOT / "figures" / "paper_v22" / "fig6_v22_known_truth_simulation.png",
        "Figure 6. Known-truth validation of the PRI-Audit evidence gate. (A) With no incremental spatial information, the strict gate made zero positive declarations in 400 replicates at each independent-patient count; whiskers are exact 95% binomial intervals. (B) With a zero-mean coordinate-specific signal beyond pseudobulk expression, power was 0.755, 0.903 and 0.808 at 12, 24 and 48 patients. (C) When source fingerprints but no disease signal were present, source-blocked patient prediction remained near chance while random-cell interpolation reached AUROC 1.000 because all test cells arose from sources seen in training. (D) When candidate and matched expression contained the same information, weakening the comparator manufactured a large apparent gain. Simulations validate the evaluation contract, not a biological claim about psoriasis.",
        "Known-truth simulation of strict-gate calibration, power, source-overlap artifact, and mismatched-reference artifact.",
    ),
}

EQUATIONS = {
    "EQUATION_1": "rᵢ,ₘ⁽⁻ᵍ⁾ = Tₘ(Xᵢ; θ̂ₘ⁽⁻ᵍ⁾),     θ̂ₘ⁽⁻ᵍ⁾ = Fit({Xⱼ : group(j) ≠ g}).     (1)",
    "EQUATION_2": "Δₘ⁽ᵈ⁾ = AUROC(y⁽ᵈ⁾, p̂candidate,m⁽ᵈ⁾) − AUROC(y⁽ᵈ⁾, p̂expression,m⁽ᵈ⁾).     (2)",
    "EQUATION_3": "Gₘ = 1[Δₘ ≥ 0.05] × 1[Lₘ,0.95 > 0] × 1[Sₘ = 1] × 1[Fₘ = 1].     (3)",
    "EQUATION_4": "Iᵢh = [nᵢ⁻¹ Σᵤ zᵢᵤh z̄ᵢ,N(u),h] / [nᵢ⁻¹ Σᵤ zᵢᵤh² + ε],     zᵢᵤh = xᵢᵤh − x̄ᵢh.     (4)",
    "EQUATION_5": "Cᵢh = nᵢ⁻¹ Σᵤ |xᵢᵤh − x̄ᵢ,N(u),h|.     (5)",
    "EQUATION_6": "P = (1 + #{Tᵇ ≥ Tobs}) / (B + 1).     (6)",
}


def set_run_font(run, size: float | None = None, *, bold=None, italic=None, color=BLACK) -> None:
    run.font.name = FONT
    props = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        props.rFonts.set(qn(f"w:{key}"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = color


def get_or_add_style(doc: Document, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def configure_styles(doc: Document) -> None:
    normal_size = 12 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 11.5
    body_after = 0 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 5
    body_line = 2.0 if BMC_PROFILE else (1.0 if REGIONSUFF_PROFILE else 1.12)
    heading_tokens = (
        [("Heading 1", 14, 10, 6), ("Heading 2", 12, 8, 4), ("Heading 3", 12, 6, 3)]
        if (REGIONSUFF_PROFILE or BMC_PROFILE)
        else [("Heading 1", 15, 12, 5), ("Heading 2", 13, 9, 4), ("Heading 3", 11.5, 7, 3)]
    )
    component_tokens = (
        [
            ("PRI Abstract", 12, WD_ALIGN_PARAGRAPH.JUSTIFY),
            ("PRI Caption", 10, WD_ALIGN_PARAGRAPH.JUSTIFY),
            ("PRI Table Caption", 10, WD_ALIGN_PARAGRAPH.LEFT),
            ("PRI Equation", 11, WD_ALIGN_PARAGRAPH.CENTER),
            ("PRI Reference", 10, WD_ALIGN_PARAGRAPH.JUSTIFY),
            ("PRI Table Text", 8.5, WD_ALIGN_PARAGRAPH.LEFT),
        ]
        if REGIONSUFF_PROFILE
        else [
            ("PRI Abstract", 12, WD_ALIGN_PARAGRAPH.JUSTIFY),
            ("PRI Caption", 10, WD_ALIGN_PARAGRAPH.JUSTIFY),
            ("PRI Table Caption", 10, WD_ALIGN_PARAGRAPH.LEFT),
            ("PRI Equation", 11, WD_ALIGN_PARAGRAPH.CENTER),
            ("PRI Reference", 10, WD_ALIGN_PARAGRAPH.JUSTIFY),
            ("PRI Table Text", 8.5, WD_ALIGN_PARAGRAPH.LEFT),
        ]
        if BMC_PROFILE
        else [
            ("PRI Abstract", 11.2, WD_ALIGN_PARAGRAPH.JUSTIFY),
            ("PRI Caption", 9.4, WD_ALIGN_PARAGRAPH.JUSTIFY),
            ("PRI Table Caption", 9.6, WD_ALIGN_PARAGRAPH.LEFT),
            ("PRI Equation", 10.8, WD_ALIGN_PARAGRAPH.CENTER),
            ("PRI Reference", 9.5, WD_ALIGN_PARAGRAPH.JUSTIFY),
            ("PRI Table Text", 7.2, WD_ALIGN_PARAGRAPH.LEFT),
        ]
    )
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(normal_size)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(body_after)
    normal.paragraph_format.line_spacing = body_line
    normal.paragraph_format.widow_control = True

    for name, size, before, after in heading_tokens:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name, size, alignment in component_tokens:
        style = get_or_add_style(doc, name)
        style.base_style = normal
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.paragraph_format.alignment = alignment
        style.paragraph_format.line_spacing = 1.0 if name == "PRI Table Text" else (2.0 if BMC_PROFILE else 1.0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(2 if name == "PRI Table Text" else (2 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 4))
    doc.styles["PRI Caption"].font.italic = True
    doc.styles["PRI Caption"].paragraph_format.keep_together = True
    doc.styles["PRI Table Caption"].font.bold = True
    doc.styles["PRI Table Caption"].paragraph_format.keep_with_next = True
    doc.styles["PRI Equation"].font.italic = True
    doc.styles["PRI Equation"].paragraph_format.space_before = Pt(3)
    doc.styles["PRI Equation"].paragraph_format.space_after = Pt(6)
    doc.styles["PRI Reference"].paragraph_format.left_indent = Inches(0.22)
    doc.styles["PRI Reference"].paragraph_format.first_line_indent = Inches(-0.22)


def clear_story(story) -> None:
    element = story._element
    for child in list(element):
        element.remove(child)
    element.append(OxmlElement("w:p"))


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1.0 if BMC_PROFILE else 1.25)
        section.right_margin = Inches(1.0 if BMC_PROFILE else 1.25)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.header_distance = Inches(0.4)
        section.footer_distance = Inches(0.4)
        clear_story(section.header)
        clear_story(section.footer)
        if BMC_PROFILE:
            add_line_numbering(section)
            footer = section.footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer.paragraph_format.space_after = Pt(0)
            run = footer.add_run()
            set_run_font(run, 10, color=BLACK)
            append_page_field(run)
        elif not REGIONSUFF_PROFILE:
            header = section.header.paragraphs[0]
            header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            header.paragraph_format.space_after = Pt(0)
            run = header.add_run("PRI-Audit | submission manuscript draft")
            set_run_font(run, 8.5, color=RGBColor(92, 92, 92))
            footer = section.footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = footer.add_run()
            set_run_font(run, 8.5, color=RGBColor(92, 92, 92))
            append_page_field(run)


def append_page_field(run) -> None:
    for tag, value in (("begin", None), ("instrText", " PAGE "), ("separate", None), ("t", "1"), ("end", None)):
        node = OxmlElement(f"w:{'fldChar' if tag in {'begin', 'separate', 'end'} else tag}")
        if tag in {"begin", "separate", "end"}:
            node.set(qn("w:fldCharType"), tag)
        else:
            node.text = value
        run._r.append(node)


def add_line_numbering(section) -> None:
    props = section._sectPr
    existing = props.find(qn("w:lnNumType"))
    if existing is not None:
        props.remove(existing)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:start"), "1")
    line_numbers.set(qn("w:restart"), "newPage")
    props.append(line_numbers)


def add_left_border(paragraph, color="AF7E00", size="18", space="7") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)
    borders.append(left)
    p_pr.append(borders)


def add_inline(paragraph, text: str, size: float) -> None:
    for piece in re.split(r"(\*\*.*?\*\*)", text):
        if not piece:
            continue
        is_bold = piece.startswith("**") and piece.endswith("**")
        content = piece[2:-2] if is_bold else piece
        run = paragraph.add_run(content)
        set_run_font(run, size, bold=is_bold)


def add_body(doc: Document, text: str, *, abstract=False) -> None:
    paragraph = doc.add_paragraph(style="PRI Abstract" if abstract else "Normal")
    add_inline(paragraph, text, 12 if (REGIONSUFF_PROFILE or BMC_PROFILE) else (11.2 if abstract else 11.5))
    return paragraph


def add_equation(doc: Document, key: str) -> None:
    paragraph = doc.add_paragraph(style="PRI Equation")
    run = paragraph.add_run(EQUATIONS[key])
    set_run_font(run, 11 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 10.8, italic=True)


def set_picture_alt_text(paragraph, title: str, description: str) -> None:
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("name", title)
        doc_pr.set("title", title)
        doc_pr.set("descr", description)


def add_figure(doc: Document, key: str) -> None:
    path, caption, alt = FIGURES[key]
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(5.64))
    set_picture_alt_text(paragraph, key.replace("_", " ").title(), alt)
    cap = doc.add_paragraph(style="PRI Caption")
    add_inline(cap, caption, 10 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 9.4)


def set_cell_shading(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width: float) -> None:
    props = cell._tc.get_or_add_tcPr()
    tc_width = props.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        props.append(tc_width)
    tc_width.set(qn("w:w"), str(int(width * 1440)))
    tc_width.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    props = table._tbl.tblPr
    layout = props.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        props.append(layout)
    layout.set(qn("w:type"), "fixed")
    total = props.first_child_found_in("w:tblW")
    if total is None:
        total = OxmlElement("w:tblW")
        props.append(total)
    total.set(qn("w:w"), str(int(sum(widths) * 1440)))
    total.set(qn("w:type"), "dxa")
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row_props = row._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_repeat(row) -> None:
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def write_cell(cell, text: str, *, header=False, center=False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = "PRI Table Text"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_together = True
    for index, piece in enumerate(str(text).split("\n")):
        if index:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        run = paragraph.add_run(piece)
        set_run_font(run, 8.5 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 7.2, bold=header)
    if header and not BMC_PROFILE:
        set_cell_shading(cell, MID_GRAY)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def add_csv_table(doc: Document, title: str, filename: str, columns: list[tuple[str, str, float, bool]], note: str) -> None:
    cap = doc.add_paragraph(style="PRI Table Caption")
    add_inline(cap, title, 10 if REGIONSUFF_PROFILE else 9.6)
    rows = read_csv(TABLES / filename)
    table = doc.add_table(rows=1, cols=len(columns))
    for cell, (display, _, _, _) in zip(table.rows[0].cells, columns):
        write_cell(cell, display, header=True, center=True)
    mark_header_repeat(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, (_, field, _, centered) in zip(cells, columns):
            if row_index % 2 and not BMC_PROFILE:
                set_cell_shading(cell, LIGHT_GRAY)
            write_cell(cell, row[field], center=centered)
    set_table_geometry(table, [width for _, _, width, _ in columns])
    note_paragraph = doc.add_paragraph(style="PRI Caption")
    add_inline(note_paragraph, note, 9.2 if REGIONSUFF_PROFILE else 8.8)


def add_table(doc: Document, key: str) -> None:
    if key == "TABLE_1":
        add_csv_table(
            doc,
            "Table 1. Public cohorts, independent units, and frozen analytical roles.",
            "table1_cohort_roles.csv",
            [
                ("Cohort", "Cohort", 0.76, True),
                ("Modality", "Modality", 0.55, False),
                ("Frozen role", "Frozen role", 1.05, False),
                ("Unit", "Independent unit", 0.50, False),
                ("n", "Raw GEO header samples", 0.32, True),
                ("PRI-Audit use", "PRI-Audit use", 2.59, False),
            ],
            "Cells and spots quantify measurement density; they never increase the independent inferential-unit count. GSE202011 is reported separately as a pre-specified external boundary replication.",
        )
    elif key == "TABLE_2":
        add_csv_table(
            doc,
            "Table 2. Strict matched representation comparisons and incremental-value boundary.",
            "table2_strict_matched_comparisons.csv",
            [
                ("Cohort", "Dataset", 0.70, True),
                ("Unit / split", "Unit / split", 0.75, False),
                ("Candidate", "Representation", 0.65, False),
                ("Matched reference", "Matched reference", 0.70, False),
                ("AUROC", "Held-out AUROC", 0.50, True),
                ("ΔAUROC (95% CI)", "Incremental AUROC (95% CI)", 1.14, True),
                ("Status", "Status", 1.33, False),
            ],
            "The primary gate requires ΔAUROC ≥ +0.05, a positive independent-unit 95% bootstrap lower bound, strict construction and a compatible counterfactual diagnostic. Saturated comparisons do not establish equivalence.",
        )
    elif key == "TABLE_3":
        add_csv_table(
            doc,
            "Table 3. Reproducibility and counterfactual audit diagnostics.",
            "table3_audit_diagnostics.csv",
            [
                ("Audit", "Audit", 1.03, False),
                ("Cohort", "Dataset", 0.88, False),
                ("Primary estimate", "Primary estimate", 1.00, False),
                ("Uncertainty / diagnostic", "Uncertainty / diagnostic", 1.48, False),
                ("Interpretation", "Interpretation", 1.38, False),
            ],
            "Diagnostic arms are shown to expose estimand mismatch or comparator artefacts. They are not alternative primary models and cannot be used to claim incremental value.",
        )


def add_title_page(doc: Document, lines: list[str]) -> int:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(6 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 10)
    title.paragraph_format.keep_with_next = True
    if not (REGIONSUFF_PROFILE or BMC_PROFILE):
        add_left_border(title)
    run = title.add_run(lines[0][2:])
    set_run_font(run, 16 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 17, bold=True, color=BLACK if (REGIONSUFF_PROFILE or BMC_PROFILE) else GOLD)
    title_lines = (
        ((lines[2], 16), (lines[4], 16), (lines[6], 16))
        if REGIONSUFF_PROFILE
        else ((lines[2], 12), (lines[4], 12), (lines[6], 12))
        if BMC_PROFILE
        else ((lines[2], 11), (lines[4], 10.2), (lines[6], 9.8))
    )
    for text, size in title_lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4 if (REGIONSUFF_PROFILE or BMC_PROFILE) else (2 if size > 10 else 8))
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(text)
        set_run_font(run, size)
    return 8


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Manuscript source not found: {SOURCE}")
    for path, _, _ in FIGURES.values():
        if not path.exists():
            raise FileNotFoundError(f"Figure missing: {path}")
    for filename in ("table1_cohort_roles.csv", "table2_strict_matched_comparisons.csv", "table3_audit_diagnostics.csv"):
        if not (TABLES / filename).exists():
            raise FileNotFoundError(f"Table missing: {TABLES / filename}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    if REFERENCE.exists():
        shutil.copy2(REFERENCE, OUTPUT)
        # The retained reference is read-only. Its content is cleared below, so
        # the working output must become a writable derived manuscript before
        # saving.  This optional path only transfers visual defaults.
        OUTPUT.chmod(0o644)
        doc = Document(OUTPUT)
    else:
        # A public release must never depend on the private visual reference.
        # All required typography, page geometry and manuscript content are
        # set below, so an empty document yields the same scientific artifact.
        doc = Document()
    clear_body(doc)
    configure_page(doc)
    configure_styles(doc)
    props = doc.core_properties
    props.title = "PRI-Audit: a falsification-first framework for testing patient-level gain from high-resolution psoriasis transcriptomics"
    props.subject = "Public-data methodological benchmark of patient-level incremental value"
    props.author = ""
    props.keywords = "psoriasis; single-cell; spatial transcriptomics; validation; reproducibility"
    props.comments = "Generated from the frozen v21 real-data package and v22 submission-closure analyses."

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = add_title_page(doc, lines)
    in_abstract = False
    in_references = False
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        token = re.fullmatch(r"\[\[((?:FIGURE|TABLE|EQUATION)_\d+)\]\]", line)
        if token:
            key = token.group(1)
            if key.startswith("FIGURE"):
                add_figure(doc, key)
            elif key.startswith("TABLE"):
                add_table(doc, key)
            else:
                add_equation(doc, key)
            continue
        if line.startswith("## "):
            heading = line[3:]
            in_abstract = heading == "Abstract"
            in_references = heading == "References"
            paragraph = doc.add_heading(heading, level=1)
            for run in paragraph.runs:
                set_run_font(run, 14 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 15, bold=True)
            continue
        if line.startswith("### "):
            paragraph = doc.add_heading(line[4:], level=2)
            for run in paragraph.runs:
                set_run_font(run, 12 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 13, bold=True)
            continue
        if in_references and re.match(r"^\d+\.\s", line):
            paragraph = doc.add_paragraph(style="PRI Reference")
            add_inline(paragraph, line, 10 if (REGIONSUFF_PROFILE or BMC_PROFILE) else 9.5)
            continue
        add_body(doc, line, abstract=in_abstract)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
