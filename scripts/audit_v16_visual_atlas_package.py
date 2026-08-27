"""Audit the final visual-atlas manuscript and its public-data figure package."""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "manuscript/PRI_Audit_psoriasis_manuscript_v3_visual_atlas.docx"
REFERENCE = ROOT / ".tmp_docx_template.87VD3w/reference.docx"
MAIN = ROOT / "figures/paper_v16/main"
ED = ROOT / "figures/paper_v16/extended_data"
TABLES = ROOT / "results/paper_v15"
PORTRAITS = ROOT / "results/paper_v16"
RENDER = ROOT / "manuscript/render_v3_visual_atlas_final"
OUT = PORTRAITS / "visual_atlas_package_audit.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def complete(stems: list[str], directory: Path) -> bool:
    return all(
        (directory / f"{stem}.{extension}").exists()
        and (directory / f"{stem}.{extension}").stat().st_size > 0
        for stem in stems
        for extension in ("png", "pdf", "svg")
    )


def main() -> None:
    main_stems = [
        "fig1_eleven_dataset_visual_atlas", "fig2_pri_audit_framework",
        "fig3_expression_ceiling", "fig4_unsupervised_optimism", "fig5_external_transfer",
        "fig6_strict_spatial_value", "fig7_falsification_evidence_map",
        "fig8_observed_data_atlas", "fig9_robustness_multiverse",
    ]
    ed_stems = [
        "figED1_cell_spot_qc", "figED2_strict_outer_roc", "figED3_prediction_distributions",
        "figED4_bulk_robustness", "figED5_neighborhood_sensitivity", "figED6_external_calibration",
        "figED7_transport_diagnostics", "figED8_spatial_graph_diagnostics",
        "figED9_complete_null_distributions", "figED10_hyperparameter_sensitivity",
        "figED11_observed_single_cell_landscapes", "figED12_predefined_spatial_gallery",
        "figED13_bulk_patient_portraits", "figED14_cytokine_perturbation_portrait",
    ]
    table_files = [
        *(f"main_table{i}_{suffix}.csv" for i, suffix in (
            (1, "public_datasets"), (2, "patient_level_incremental_value"),
            (3, "resolution_specific_falsification"), (4, "diagnostics_and_sensitivity"),
        )),
        "supp_table1_high_resolution_qc.csv", "supp_table2_strict_resampling.csv",
        "supp_table3_full_incremental_evidence.csv", "supp_table4_falsification.csv",
        "supp_table5_external_transfer_calibration.csv", "supp_table6_spatial_sensitivity_summary.csv",
        "supp_table7_claim_registry.csv",
    ]
    portrait_files = [
        "dataset_portrait_manifest.json", "bulk_patient_portrait_values.csv",
        "gse24767_mean_portrait_values.csv", "gse24767_donor_portrait_values.csv",
        "a11y_audit.json",
    ]

    document = Document(DOCX)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    main_captions = [p.text for p in document.paragraphs if re.match(r"^Figure \d+\.", p.text)]
    ed_headings = [
        p.text for p in document.paragraphs
        if p.text.startswith("Extended Data Figure") and p.style.name == "Heading 2"
    ]
    with zipfile.ZipFile(DOCX) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        names = set(archive.namelist())
    tracked = bool(re.search(r"<w:(ins|del|moveFrom|moveTo)\b", document_xml))
    comments = "word/comments.xml" in names or "word/commentsExtended.xml" in names
    page_pngs = sorted(RENDER.glob("page-*.png"))
    dataset_ids = [
        "GSE117239", "GSE136757", "GSE153007", "GSE182740", "GSE51440", "GSE57225",
        "GSE85034", "GSE151177", "GSE173706", "GSE206391", "GSE24767",
    ]
    audit = {
        "pass": True,
        "main_figures": {"count": 9, "formats_each": 3, "complete": complete(main_stems, MAIN)},
        "extended_data_figures": {"count": 14, "formats_each": 3, "complete": complete(ed_stems, ED)},
        "tables": {"count": 11, "complete": all((TABLES / name).exists() for name in table_files)},
        "portrait_source_exports": {
            "count": len(portrait_files),
            "complete": all((PORTRAITS / name).exists() for name in portrait_files),
        },
        "docx": {
            "exists": DOCX.exists(),
            "bytes": DOCX.stat().st_size,
            "sha256": sha256(DOCX),
            "inline_images": len(document.inline_shapes),
            "tables": len(document.tables),
            "main_caption_numbers": [int(re.match(r"^Figure (\d+)\.", item).group(1)) for item in main_captions],
            "extended_heading_count": len(ed_headings),
            "all_dataset_ids_present": all(dataset in text for dataset in dataset_ids),
            "unresolved_tokens": re.findall(r"\[\[[A-Z0-9_]+\]\]", text),
            "tracked_changes": tracked,
            "comments_part": comments,
        },
        "render": {
            "pages": len(page_pngs),
            "pdf_exists": (RENDER / "PRI_Audit_psoriasis_manuscript_v3_visual_atlas.pdf").exists(),
            "all_pages_nonempty": all(path.stat().st_size > 0 for path in page_pngs),
        },
        "reference_sha256": sha256(REFERENCE),
        "reference_matches_frozen_hash": sha256(REFERENCE)
        == "68932101fe50e45b07bc49d9c928b373d10866e7a13ad0c6c0a618056aa6a5e3",
    }
    audit["pass"] = all([
        audit["main_figures"]["complete"], audit["extended_data_figures"]["complete"],
        audit["tables"]["complete"], audit["portrait_source_exports"]["complete"],
        audit["docx"]["inline_images"] == 23, audit["docx"]["tables"] == 11,
        audit["docx"]["main_caption_numbers"] == list(range(1, 10)),
        audit["docx"]["extended_heading_count"] == 14,
        audit["docx"]["all_dataset_ids_present"], not audit["docx"]["unresolved_tokens"],
        not tracked, not comments, audit["render"]["pages"] == 55,
        audit["render"]["pdf_exists"], audit["render"]["all_pages_nonempty"],
        audit["reference_matches_frozen_hash"],
    ])
    OUT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(json.dumps(audit, indent=2))
    if not audit["pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
